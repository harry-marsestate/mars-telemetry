"""Parses the three real labour source files into labour_actuals rows.

  - Silverado_2023_Hours_Analysis.xlsx 'Raw Data'  -> vintage 2023
  - Silverado_2024_Hours_Analysis.xlsx 'Raw Data'  -> vintage 2024
  - Mars_Invoice_Backup_7_31_26.xlsx   'Labor'      -> vintage 2026 (July)
  - Mars_Invoice_Backup_7_31_26.xlsx   'Expenses'   -> vintage 2026 (July)

The Mars Invoice Backup file is JULY 2026 -- confirmed by the filename
(7_31_26), every Expenses row's own date (2026-07-xx), and the user
directly. It is NOT 2025 data.

Every parse function returns (rows, checksum). checksum compares computed
totals against CHECKSUMS below -- values independently verified against
each file's own TOTAL row before this module was written (see
docs/SECURITY.md), then re-confirmed directly against the live files in
this session (openpyxl/pandas read matched every one exactly except the
2024 file's own narrative claim of "24 monthly Silverado invoices" on its
Summary tab, which does not match the 20 distinct invoice numbers actually
present in its Raw Data tab -- flagged, not silently reconciled; every
numeric total, including the Farming/Development split, matches exactly).
backfill.py refuses to insert a source whose checksum doesn't reconcile.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

RAW_DIR = Path(__file__).parent.parent.parent / "seed-data" / "labour_raw"

_CODE_RE = re.compile(r"^(\d+(?:\.\d+)?)\s+(.*)$")

# Expense account -> folded job_category. Strips the leading account
# number and maps by name -- 'Toilet, Other' is a genuinely new category
# ('Other'), not folded into an existing labor category.
ACCOUNT_CATEGORY_MAP = {
    "Fertilize": "Fertilize",
    "Disease Control": "Disease Control",
    "Irrigation": "Irrigation",
    "Toilet, Other": "Other",
}

CHECKSUMS = {
    "2023": {"rows": 313, "hours": 3915.989, "amount": 205604.13, "job_categories": 18, "invoices": 8},
    "2024": {
        "rows": 481, "hours": 5563.12, "amount": 298652.13, "job_categories": 19,
        "farming_hours": 3112.69, "farming_amount": 173719.72,
        "development_hours": 2450.43, "development_amount": 124932.41,
    },
    "mars_invoice_labor": {
        "rows": 65, "hours": 581.24, "amount": 29331.39, "job_categories": 12,
        "development_hours": 393.2, "development_amount": 19102.67,
        "farming_hours": 188.04, "farming_amount": 10228.72,
    },
    "mars_invoice_expenses": {
        "rows": 20, "amount": 4346.0928,
        "by_category": {"Disease Control": 3217.3568, "Fertilize": 656.9472, "Irrigation": 231.84, "Other": 239.9488},
    },
}
_TOL = 0.01


def _split_code(raw: str) -> tuple[str | None, str]:
    """'028.00 General Labor' -> ('028.00', 'General Labor'); no code prefix -> (None, raw)."""
    m = _CODE_RE.match(raw.strip())
    return (m.group(1), m.group(2)) if m else (None, raw.strip())


def _month_start(month_name: str, year: int) -> str:
    return pd.Timestamp(f"{month_name} 1, {year}").date().isoformat()


def parse_2023(path: Path = RAW_DIR / "Silverado_2023_Hours_Analysis.xlsx") -> tuple[list[dict], dict]:
    df = pd.read_excel(path, sheet_name="Raw Data")
    df = df[df["Month"].notna()].reset_index(drop=True)

    rows = []
    for i, r in df.iterrows():
        task_code, task = _split_code(str(r["Task"]))
        role_code, role = _split_code(str(r["Role"]))
        rows.append({
            "vintage": 2023,
            "period_month": _month_start(r["Month"], 2023),
            "invoice_number": str(int(r["Invoice #"])),
            "invoice_type": None,
            "job_category": r["Job Category"],
            "task": task, "task_code": task_code,
            "role": role, "role_code": role_code,
            "hours": float(r["Hours"]), "rate_per_hour": float(r["Rate ($/hr)"]),
            "amount_usd": float(r["Amount ($)"]),
            "entry_kind": "labor",
            "expense_vendor": None, "expense_memo": None, "expense_account": None, "expense_date": None,
            "source_system": "silverado_hours_analysis",
            "source_file": "Silverado_2023_Hours_Analysis.xlsx",
            "source_row_id": i + 2,  # +2: 1-based, plus header row
        })

    c = CHECKSUMS["2023"]
    check = {
        "label": "2023 Raw Data",
        "checks": [
            ("rows", len(rows), c["rows"]),
            ("hours", sum(r["hours"] for r in rows), c["hours"]),
            ("amount", sum(r["amount_usd"] for r in rows), c["amount"]),
            ("job_categories", df["Job Category"].nunique(), c["job_categories"]),
            ("invoices", df["Invoice #"].nunique(), c["invoices"]),
        ],
    }
    return rows, check


def parse_2024(path: Path = RAW_DIR / "Silverado_2024_Hours_Analysis.xlsx") -> tuple[list[dict], dict]:
    df = pd.read_excel(path, sheet_name="Raw Data")
    df = df[df["Month"].notna()].reset_index(drop=True)

    rows = []
    for i, r in df.iterrows():
        task_code, task = _split_code(str(r["Task"]))
        role_code, role = _split_code(str(r["Role"]))
        rows.append({
            "vintage": 2024,
            "period_month": _month_start(r["Month"], 2024),
            "invoice_number": str(int(r["Invoice #"])),
            "invoice_type": r["Invoice Type"],
            "job_category": r["Job Category"],
            "task": task, "task_code": task_code,
            "role": role, "role_code": role_code,
            "hours": float(r["Hours"]), "rate_per_hour": float(r["Rate ($/hr)"]),
            "amount_usd": float(r["Amount ($)"]),
            "entry_kind": "labor",
            "expense_vendor": None, "expense_memo": None, "expense_account": None, "expense_date": None,
            "source_system": "silverado_hours_analysis",
            "source_file": "Silverado_2024_Hours_Analysis.xlsx",
            "source_row_id": i + 2,
        })

    c = CHECKSUMS["2024"]
    farming = df[df["Invoice Type"] == "Farming"]
    development = df[df["Invoice Type"] == "Development"]
    check = {
        # NOTE: this file's own Summary tab claims "24 monthly Silverado
        # invoices" in its narrative description, but Raw Data contains
        # only 20 distinct Invoice # values -- a discrepancy in the
        # source file's own narrative text, not in this parse. Every
        # numeric total (hours/amount/split) matches exactly, so this is
        # reported as a non-blocking flag, not a refusal cause.
        "label": "2024 Raw Data",
        "checks": [
            ("rows", len(rows), c["rows"]),
            ("hours", sum(r["hours"] for r in rows), c["hours"]),
            ("amount", sum(r["amount_usd"] for r in rows), c["amount"]),
            ("job_categories", df["Job Category"].nunique(), c["job_categories"]),
            ("farming_hours", farming["Hours"].sum(), c["farming_hours"]),
            ("farming_amount", farming["Amount ($)"].sum(), c["farming_amount"]),
            ("development_hours", development["Hours"].sum(), c["development_hours"]),
            ("development_amount", development["Amount ($)"].sum(), c["development_amount"]),
        ],
        "flags": [
            f"Summary tab claims 24 invoices; Raw Data has {df['Invoice #'].nunique()} distinct "
            f"Invoice # values ({sorted(int(x) for x in df['Invoice #'].unique())}). "
            f"Not blocking: no numeric checksum depends on invoice count."
        ],
    }
    return rows, check


def parse_mars_invoice_labor(path: Path = RAW_DIR / "Mars_Invoice_Backup_7_31_26.xlsx") -> tuple[list[dict], dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Labor"]

    # Invoice numbers taken from the two real PDF invoices backing this
    # file (Mars Development INV 39517.pdf / Mars Farming INV 39518.pdf,
    # seed-data/labour_raw/ -- same two divisions, same July 2026 period).
    INVOICE_BY_DIVISION = {"Mars Development": "39517", "Mars Farming": "39518"}

    rows = []
    division = None
    for r in range(2, ws.max_row + 1):
        b, c_ = ws.cell(r, 2).value, ws.cell(r, 3).value
        d, e, f = ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value
        a = ws.cell(r, 1).value
        if a == "TOTAL":
            continue
        if b is not None:
            if isinstance(b, str) and b.startswith("Total "):
                continue
            division = b.split(":")[-1]  # 'Mars Estates Inc.:Mars Development' -> 'Mars Development'
            continue
        if c_ is None or ":" not in str(c_):
            continue
        job_category, task_raw, role_raw = str(c_).split(":", 2)
        task_code, task = _split_code(task_raw)
        role_code, role = _split_code(role_raw)
        rows.append({
            "vintage": 2026,
            "period_month": "2026-07-01",
            "invoice_number": INVOICE_BY_DIVISION[division],
            "invoice_type": division,
            "job_category": job_category,
            "task": task, "task_code": task_code,
            "role": role, "role_code": role_code,
            "hours": float(d), "rate_per_hour": float(e),
            "amount_usd": float(f),
            "entry_kind": "labor",
            "expense_vendor": None, "expense_memo": None, "expense_account": None, "expense_date": None,
            "source_system": "mars_invoice_labor",
            "source_file": "Mars_Invoice_Backup_7_31_26.xlsx:Labor",
            "source_row_id": r,
        })

    c = CHECKSUMS["mars_invoice_labor"]
    dev = [r for r in rows if r["invoice_type"] == "Mars Development"]
    farm = [r for r in rows if r["invoice_type"] == "Mars Farming"]
    check = {
        "label": "Mars Invoice Backup: Labor",
        "checks": [
            ("rows", len(rows), c["rows"]),
            ("hours", sum(r["hours"] for r in rows), c["hours"]),
            ("amount", sum(r["amount_usd"] for r in rows), c["amount"]),
            ("job_categories", len({r["job_category"] for r in rows}), c["job_categories"]),
            ("development_hours", sum(r["hours"] for r in dev), c["development_hours"]),
            ("development_amount", sum(r["amount_usd"] for r in dev), c["development_amount"]),
            ("farming_hours", sum(r["hours"] for r in farm), c["farming_hours"]),
            ("farming_amount", sum(r["amount_usd"] for r in farm), c["farming_amount"]),
        ],
    }
    return rows, check


def parse_mars_invoice_expenses(path: Path = RAW_DIR / "Mars_Invoice_Backup_7_31_26.xlsx") -> tuple[list[dict], dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Expenses"]

    rows = []
    division = None
    for r in range(2, ws.max_row + 1):  # row 1 is the header
        a = ws.cell(r, 1).value
        c_ = ws.cell(r, 3).value
        dtype = ws.cell(r, 4).value
        if a == "TOTAL":
            continue
        if c_ is not None:
            if isinstance(c_, str) and c_.startswith("Total "):
                continue
            division = c_  # 'Mars Development' / 'Mars Farming'
            continue
        if dtype is None:
            continue
        date_val = ws.cell(r, 5).value
        source_name = ws.cell(r, 6).value
        memo = ws.cell(r, 7).value
        account = ws.cell(r, 8).value
        amount = ws.cell(r, 9).value
        account_name = account.split("·")[-1].strip() if account else None
        job_category = ACCOUNT_CATEGORY_MAP.get(account_name, account_name)
        rows.append({
            "vintage": 2026,
            "period_month": date_val.replace(day=1).date().isoformat(),
            "invoice_number": None,
            "invoice_type": division,
            "job_category": job_category,
            "task": None, "task_code": None, "role": None, "role_code": None,
            "hours": None, "rate_per_hour": None,
            "amount_usd": float(amount),
            "entry_kind": "expense",
            "expense_vendor": source_name, "expense_memo": memo,
            "expense_account": account, "expense_date": date_val.date().isoformat(),
            "source_system": "mars_invoice_expense",
            "source_file": "Mars_Invoice_Backup_7_31_26.xlsx:Expenses",
            "source_row_id": r,
        })

    c = CHECKSUMS["mars_invoice_expenses"]
    by_cat: dict[str, float] = {}
    for r in rows:
        by_cat[r["job_category"]] = by_cat.get(r["job_category"], 0.0) + r["amount_usd"]

    checks = [
        ("rows", len(rows), c["rows"]),
        ("amount", sum(r["amount_usd"] for r in rows), c["amount"]),
    ]
    for cat, target in c["by_category"].items():
        checks.append((f"category:{cat}", by_cat.get(cat, 0.0), target))

    return rows, {"label": "Mars Invoice Backup: Expenses", "checks": checks}


def evaluate(check: dict) -> bool:
    """Prints every check line and returns True iff all reconcile within tolerance."""
    ok = True
    print(f"  self-check ({check['label']}, vs. verified source-file checksums):")
    for name, computed, target in check["checks"]:
        diff = computed - target
        passed = abs(diff) <= _TOL
        ok = ok and passed
        print(f"    {name}: computed={computed}, target={target}, diff={diff:+.4f}  [{'OK' if passed else 'FAIL'}]")
    for flag in check.get("flags", []):
        print(f"  FLAG (non-blocking): {flag}")
    return ok
